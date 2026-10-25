from io import BytesIO

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.api import documents
from backend.config import settings
from backend.main import app
from backend.storage.document_store import DocumentStore


@pytest.fixture
def store(tmp_path, monkeypatch):
    document_store = DocumentStore(tmp_path / "uploads", ttl_minutes=30)
    app.dependency_overrides[documents.get_document_store] = lambda: document_store
    monkeypatch.setattr(settings, "max_upload_size_mb", 10)
    yield document_store
    app.dependency_overrides.clear()


@pytest.fixture
def client(store):
    return TestClient(app)


def make_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Quarterly report: revenue increased.")
    content = document.tobytes()
    document.close()
    return content


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Introductory paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "42"
    document.add_paragraph("Closing paragraph.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def make_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Annual Results"
    sheet.append(["Year", "Revenue"])
    sheet.append([2025, 42])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize(
    ("filename", "content", "file_type"),
    [
        ("report.pdf", make_pdf(), "pdf"),
        ("report.docx", make_docx(), "docx"),
        ("notes.txt", "A plain text note.".encode(), "txt"),
        ("results.csv", b"year,revenue\n2025,42\n", "csv"),
        ("results.xlsx", make_xlsx(), "xlsx"),
    ],
    ids=["pdf", "docx", "text", "csv", "xlsx"],
)
def test_upload_extracts_supported_formats(client, filename, content, file_type):
    response = client.post("/api/v1/documents/upload", files={"file": (filename, content)})

    assert response.status_code == 201
    payload = response.json()
    assert payload["document_id"]
    assert payload["detected_file_type"] == file_type
    assert payload["status"] == "completed"
    assert payload["file_size"] == len(content)
    assert "extracted_text" in payload
    assert "source_references" in payload
    if file_type in {"csv", "xlsx", "docx"}:
        assert payload["tables"]
    if file_type == "pdf":
        assert {"type": "page", "page": 1} in payload["source_references"]
    if file_type == "xlsx":
        assert payload["tables"][0]["name"] == "Annual Results"


def test_docx_preserves_paragraph_and_table_content(client):
    response = client.post("/api/v1/documents/upload", files={"file": ("ordered.docx", make_docx())})

    assert response.status_code == 201
    text = response.json()["extracted_text"]
    assert text.index("Introductory paragraph") < text.index("[Table 1]") < text.index("Closing paragraph")
    assert "Revenue | 42" in text


def test_unsupported_extension_is_rejected(client):
    response = client.post("/api/v1/documents/upload", files={"file": ("payload.exe", b"not executable")})

    assert response.status_code == 415


@pytest.mark.parametrize(
    ("filename", "content"),
    [("empty.txt", b""), ("broken.pdf", b"%PDF-1.7 not a real PDF")],
)
def test_empty_and_corrupt_documents_are_rejected(client, filename, content):
    response = client.post("/api/v1/documents/upload", files={"file": (filename, content)})

    assert response.status_code == 422


def test_oversized_upload_is_rejected(client, monkeypatch):
    monkeypatch.setattr(settings, "max_upload_size_mb", 1)
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("large.txt", b"a" * (1024 * 1024 + 1))},
    )

    assert response.status_code == 413


def test_parser_error_returns_clear_error_and_removes_temp_file(client, store, monkeypatch):
    def fail_parser(*_args):
        raise ValueError("Parser could not read the document.")

    monkeypatch.setattr(documents, "parse_document", fail_parser)
    response = client.post("/api/v1/documents/upload", files={"file": ("bad.txt", b"some text")})

    assert response.status_code == 422
    assert "could not read" in response.json()["detail"]
    assert list(store.directory.iterdir()) == []


def test_results_status_and_unknown_document_ids(client):
    upload = client.post("/api/v1/documents/upload", files={"file": ("notes.txt", b"hello source")})
    document_id = upload.json()["document_id"]

    status = client.get(f"/api/v1/documents/{document_id}/status")
    results = client.get(f"/api/v1/documents/{document_id}/results")
    missing = client.get("/api/v1/documents/not-a-real-id/results")

    assert status.status_code == 200
    assert status.json()["status"] == "completed"
    assert results.status_code == 200
    assert results.json()["extracted_text"] == "hello source"
    assert missing.status_code == 404


def test_delete_removes_temporary_file(client, store):
    upload = client.post("/api/v1/documents/upload", files={"file": ("remove.txt", b"temporary")})
    document_id = upload.json()["document_id"]
    temp_path = store.get(document_id).temp_path
    assert temp_path.exists()

    deleted = client.delete(f"/api/v1/documents/{document_id}")

    assert deleted.status_code == 200
    assert deleted.json() == {"document_id": document_id, "status": "deleted"}
    assert not temp_path.exists()
    assert client.get(f"/api/v1/documents/{document_id}/results").status_code == 404


def test_expired_document_is_removed_on_next_store_access(client, store):
    upload = client.post("/api/v1/documents/upload", files={"file": ("expire.txt", b"temporary")})
    document_id = upload.json()["document_id"]
    temp_path = store.get(document_id).temp_path
    store.get(document_id).expires_at = 0

    response = client.get(f"/api/v1/documents/{document_id}/status")

    assert response.status_code == 404
    assert not temp_path.exists()


def test_path_like_filename_is_safely_reduced(client, store):
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": (r"..\..\private\report.txt", b"safe content")},
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["original_filename"] == "report.txt"
    stored = store.get(payload["document_id"])
    assert stored.temp_path.parent == store.directory
    assert stored.temp_path.name != payload["original_filename"]
